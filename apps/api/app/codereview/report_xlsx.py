"""Code Security Review XLSX tracker builder (module reference §5).

Numbers/strings come ONLY from `review.report` (the ingest-normalized
JSONB, module reference §3) — never recomputed, never network/LLM. Every
attacker-controlled string (it originated in the scanned repo, passed
through an LLM) goes through `_guard` before it reaches a cell.

Presentation rules (user feedback 2026-09-12): professional headings,
numbers centred, tinted colour for severity / verdict / status, long prose
split into bullet points with code tokens, risk words and fix words
highlighted (rich text), and an Exploit Chains sheet that explains itself.
"""

import io
import json
import re
from datetime import datetime
from pathlib import Path

from app.mitre.report_common import resolve_branding

BRAND = "341954"
ACCENT = "00A98B"
ZEBRA = "F3F0F7"
SECTION = "E9E4F0"
INK = "1E293B"
MUTED = "64748B"

# (fill, text) — same tints as the web UI (bg-{c}-100 / text-{c}-800).
_SEVERITY_FILLS = {
    "critical": ("FFE4E6", "9F1239"),
    "high": ("FFEDD5", "9A3412"),
    "medium": ("FEF3C7", "92400E"),
    "low": ("D1FAE5", "065F46"),
    "info": ("E0F2FE", "075985"),
}
_VERDICT = {
    "TRUE_POSITIVE": ("Confirmed", "D1FAE5", "065F46"),
    "FALSE_POSITIVE": ("False positive", "F1F5F9", "475569"),
}
_STATUS_OPTIONS = ["Open", "In progress", "Fixed", "Accepted risk", "False positive"]
_STATUS_FILLS = {
    "Open": ("FFE4E6", "9F1239"),
    "In progress": ("FEF3C7", "92400E"),
    "Fixed": ("D1FAE5", "065F46"),
    "Accepted risk": ("E0F2FE", "075985"),
    "False positive": ("F1F5F9", "475569"),
}

# Run-extras display labels (contract §"Display labels", shared by XLSX/PPTX/web).
FIX_STATUS_LABELS = {
    "fixed": "Fixed (scanner accepted)",
    "patch_rejected": "Patch rejected by scanner policy",
    "needs_review": "Not fixed: needs manual review",
    "not_fixed": "Not fixed",
    "not_attempted": "Not attempted (report-only run)",
}
FIX_TESTS_LABELS = {
    "broke_tests": "Fix broke a test",
    "passed": "Tests passed",
    "not_tested": "No test results for this code",
    "no_test_results": "No test results uploaded",
}

# Highlighting: single source of truth is highlight_words.json (next to this
# file); the web drawer's highlightWords.ts is GENERATED from it by
# scripts/generate_highlight_words.py, and test_codereview_report.py fails
# when that mirror is stale — so XLSX, PPTX and UI can no longer drift.
_HIGHLIGHT_WORDS = json.loads(
    (Path(__file__).with_name("highlight_words.json")).read_text(encoding="utf-8")
)
_RISK_RE = re.compile(r"\b(" + "|".join(_HIGHLIGHT_WORDS["risk"]) + r")\b", re.I)
_FIX_RE = re.compile(r"\b(" + "|".join(_HIGHLIGHT_WORDS["fix"]) + r")\b", re.I)
_CODE_RE = re.compile(
    r"`[^`]+`|https?://[^\s)]+|\b(?:GET|POST|PUT|PATCH|DELETE)\s+/[\w\-./:?=&{}]*|"
    r"\b(?:[A-Za-z_$][\w$]*\.)+[A-Za-z_$][\w$]*(?:\(\))?|\b[a-z][\w$]*[A-Z][\w$]*(?:\(\))?|"
    r"\b[\w.-]+\.(?:js|ts|py|json|yml|yaml|html|env)\b|\b\w+\(\)|/(?:[\w\-]+/)+[\w\-.]*"
)
_ABBREV_RE = re.compile(r"^(e\.g|i\.e|etc|vs|cf)$", re.I)
# (?<!\.\.) keeps an ellipsis ("SELECT ... FOR UPDATE") from ending a sentence.
_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])(?<!\.\.)\s+(?=[A-Z`(])")


def _guard(value):
    """Excel formula-injection guard: a leading =, +, - or @ in an
    attacker-controlled string would execute as a formula when the
    register opens in Excel."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _sentences(text: str) -> list:
    text = re.sub(r"\s+", " ", text or "").strip()
    if not text:
        return []
    protected = re.sub(r"\b(e\.g|i\.e|etc|vs|cf)\.\s", lambda m: m.group(1) + "․ ", text, flags=re.I)
    return [s.replace("․", ".").strip() for s in _SENTENCE_SPLIT.split(protected) if s.strip()]


def _rich(text, fix=False, bullets=True):
    """Rich-text cell: bullets per sentence, code tokens in mono, risk words
    in rose (or fix words in green when `fix`). Falls back to a plain
    guarded string when there is nothing to highlight."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    normal = InlineFont(rFont="Calibri", sz=10, color=INK)
    code = InlineFont(rFont="Consolas", sz=9.5, color="1D4ED8")
    hot = InlineFont(rFont="Calibri", sz=10, b=True, color="0F766E" if fix else "BE123C")
    word_re = _FIX_RE if fix else _RISK_RE

    parts = _sentences(text) if bullets else ([text] if text else [])
    if not parts:
        return ""
    blocks = []
    any_highlight = False
    for i, sentence in enumerate(parts):
        prefix = ("• " if bullets and len(parts) > 1 else "") if i == 0 else "\n• "
        if prefix:
            blocks.append(TextBlock(normal, prefix))
        pos = 0
        for m in _CODE_RE.finditer(sentence):
            if m.start() > pos:
                any_highlight |= _words(blocks, sentence[pos:m.start()], word_re, normal, hot)
            tok = m.group(0)
            if _ABBREV_RE.match(tok):
                any_highlight |= _words(blocks, tok, word_re, normal, hot)
            else:
                blocks.append(TextBlock(code, tok.strip("`")))
                any_highlight = True
            pos = m.end()
        if pos < len(sentence):
            any_highlight |= _words(blocks, sentence[pos:], word_re, normal, hot)
    if not any_highlight and len(parts) == 1:
        return _guard(parts[0])
    if not any_highlight:
        return _guard("\n".join("• " + p for p in parts))
    return CellRichText(*blocks)


def _words(blocks, chunk, word_re, normal, hot) -> bool:
    from openpyxl.cell.rich_text import TextBlock

    hit = False
    pos = 0
    for m in word_re.finditer(chunk):
        if m.start() > pos:
            blocks.append(TextBlock(normal, chunk[pos:m.start()]))
        blocks.append(TextBlock(hot, m.group(0)))
        hit = True
        pos = m.end()
    if pos < len(chunk):
        blocks.append(TextBlock(normal, chunk[pos:]))
    return hit


def _plain_len(value) -> int:
    if value is None:
        return 0
    if isinstance(value, str):
        return len(value)
    try:  # CellRichText
        return sum(len(str(getattr(b, "text", b))) for b in value)
    except TypeError:
        return len(str(value))


def build_xlsx_export(review) -> bytes:
    branding = resolve_branding(None)
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    report = review.report or {}
    counts = report.get("counts") or {}
    findings = report.get("findings") or []
    chains = report.get("chains") or []
    metrics = report.get("metrics") or {}
    manifest = report.get("manifest") or {}
    title_by_idx = {f.get("idx"): f.get("title") for f in findings}
    run_extras = report.get("run_extras") or None
    fix_findings = [f for f in findings if (f.get("fix") or {}).get("status") not in (None, "not_attempted")]

    wb = Workbook()
    thin = Side(style="thin", color="D6D0E0")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color="FFFFFF", size=10.5)
    body_font = Font(size=10, color=INK)
    bold_ink = Font(bold=True, size=10, color=INK)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    def fill(color):
        return PatternFill(start_color=color, end_color=color, fill_type="solid")

    def chip(cell, colors, text=None, bold=True):
        bg, fg = colors
        cell.fill = fill(bg)
        cell.font = Font(bold=bold, size=10, color=fg)
        cell.alignment = center
        if text is not None:
            cell.value = text

    def sheet(title, headers, rows, widths, first=False, filters=True, borders=True, tab=None):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        if tab:
            ws.sheet_properties.tabColor = tab
        ws.append([_guard(h) for h in headers])
        ws.row_dimensions[1].height = 24
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = fill(BRAND)
            cell.alignment = center
            if borders:
                cell.border = cell_border
        for row in rows:
            ws.append([v if not isinstance(v, str) else _guard(v) for v in row])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = top_left
                cell.font = body_font
                if borders:
                    cell.border = cell_border
                if cell.row % 2 == 0:
                    cell.fill = fill(ZEBRA)
        ws.freeze_panes = "A2"
        if filters and rows:
            ws.auto_filter.ref = ws.dimensions
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.sheet_view.showGridLines = False
        return ws

    # ------------------------------------------------------------- Read Me
    readme_rows = [
        ["What this workbook is",
         "A tracker for the findings of one code security scan, produced by Visa's "
         "open-source Vulnerability Agentic Harness (VVAH) and formatted by ScopeSense."],
        ["Findings are candidates, not confirmed bugs",
         "Every row was generated by an AI scanner. A person must confirm each one "
         "before acting on it."],
        ["Sheets",
         "Summary - the numbers at a glance.  Findings Register - one row per finding, "
         "in plain language, with tracker columns for your team.  Exploit Chains - "
         "findings an attacker would combine, in attack order."],
        ["Severity",
         "The scanner's triage rating: Critical > High > Medium > Low > Info. "
         "Work the Critical and High rows first."],
        ["CVSS / Rating",
         "Standard 0-10 severity score and its band, when the scanner computed one."],
        ["Confidence / Votes",
         "How sure the scanner is (0-100%) and how many of its internal reviewers agreed."],
        ["Verdict",
         "Confirmed = the scanner's second-pass verifier agreed it is real. "
         "False positive = the verifier rejected it. Blank = not verified."],
        ["Tracker columns",
         "Owner, Status (dropdown), Target date and Notes are yours to fill in."],
        ["Colours",
         "Rose = critical / open, orange = high, amber = medium / in progress, "
         "green = low / fixed / confirmed, blue = info / accepted risk, grey = false positive. "
         "In text: blue mono = code, file or endpoint; rose bold = attack term; green bold = fix action."],
        ["Attribution",
         "Findings produced by Visa Vulnerability Agentic Harness (Apache-2.0). "
         "ScopeSense consumes its output only; no VVAH code is vendored."],
    ]
    if run_extras:
        readme_rows.append([
            "Scanner run extras",
            "This upload included the scanner's full run folder, so some findings also carry a deep "
            "analysis and (in fix mode) an attempted patch with test results — see the Deep-verified, "
            "Fix status and Tests after fix columns in the Findings Register, the Fixes sheet (attempted "
            "fixes) and the Scan Coverage sheet (coverage, health and test-build detail).",
        ])
    ws_rm = sheet("Read Me", ["How to read this workbook", ""], readme_rows,
                  [38, 100], first=True, filters=False, borders=False, tab=BRAND)
    for row in ws_rm.iter_rows(min_row=2, max_col=1):
        row[0].font = bold_ink

    # -------------------------------------------------------------- Summary
    by_sev = counts.get("by_severity") or {}
    by_class = counts.get("by_vuln_class") or {}
    sum_rows = [["Counts by severity", ""]]
    for sev in ("critical", "high", "medium", "low", "info"):
        sum_rows.append([sev.title(), by_sev.get(sev, 0)])
    sum_rows.append(["Total findings", counts.get("total", len(findings))])
    if report.get("dropped_count"):
        sum_rows.append(["Not shown (over the 2000-finding cap)", report.get("dropped_count")])
    sum_rows.append([])
    sum_rows.append(["Counts by vulnerability type", ""])
    for label, n in sorted(by_class.items(), key=lambda kv: -kv[1]):
        sum_rows.append([label, n])
    if metrics:
        sum_rows.append([])
        sum_rows.append(["Scan metrics", ""])
        for label, key in (
            ("Scan duration (minutes)", "duration_sec"),
            ("Files in scope", "total_files_in_scope"),
            ("Files analysed", "analyzed_files_unique"),
            ("Confirmed by verifier", "true_positive_count"),
            ("Rejected by verifier (false positives)", "false_positive_count"),
            ("Tokens used", "total_tokens"),
        ):
            if metrics.get(key) is not None:
                val = metrics[key]
                if key == "duration_sec":
                    val = round(val / 60, 1)
                sum_rows.append([label, val])
    if manifest:
        sum_rows.append([])
        sum_rows.append(["Scan run details", ""])
        if manifest.get("target_git_sha"):
            sum_rows.append(["Git commit", manifest["target_git_sha"]])
        if manifest.get("tool_version") or report.get("tool_version"):
            sum_rows.append(["Scanner version", manifest.get("tool_version") or report.get("tool_version")])
        for role, model in (manifest.get("models") or {}).items():
            sum_rows.append([f"Model - {role}", model.get("id")])
        if manifest.get("total_cost_usd") is not None:
            sum_rows.append(["Total cost (USD)", manifest["total_cost_usd"]])
    if run_extras:
        sum_rows.append([])
        sum_rows.append(["Scanner run extras", ""])
        sum_rows.append(["Deep-verified findings", f"{run_extras.get('deep_verified', 0)} of {run_extras.get('total', len(findings))}"])
        for status, n in (run_extras.get("fix_counts") or {}).items():
            sum_rows.append([FIX_STATUS_LABELS.get(status, status), n])
        tests = run_extras.get("tests") or {}
        if tests.get("build"):
            sum_rows.append(["Test build status", tests["build"].title()])
    ws_sum = sheet("Summary", ["Metric", "Value"], sum_rows, [42, 26], filters=False, tab=ACCENT)
    for row in ws_sum.iter_rows(min_row=2, max_col=2):
        label, value = row
        key = str(label.value or "")
        if key in ("Counts by severity", "Counts by vulnerability type", "Scan metrics", "Scan run details", "Scanner run extras"):
            for c in row:
                c.fill = fill(SECTION)
                c.font = Font(bold=True, size=10.5, color=BRAND)
        elif key.lower() in _SEVERITY_FILLS:
            chip(label, _SEVERITY_FILLS[key.lower()])
            label.alignment = Alignment(horizontal="left", vertical="center")
        if isinstance(value.value, (int, float)):
            value.alignment = center
            value.number_format = "#,##0.0" if isinstance(value.value, float) else "#,##0"
            if key == "Total findings":
                value.font = bold_ink
                label.font = bold_ink
        elif value.value:
            value.alignment = center
            if key in ("Git commit",) or key.startswith("Model"):
                value.font = Font(name="Consolas", size=9.5, color=INK)

    # -------------------------------------------------------- Findings Register
    reg_headers = [
        "#", "Severity", "Title", "Type", "CWE", "CVSS", "Rating",
        "Confidence", "Votes", "Verdict", "File", "Lines",
        "What is wrong", "Why it matters", "How to fix", "How it is exploited",
        "Preconditions", "Code", "Source → Sink", "Also at",
        "Owner", "Status", "Target date", "Notes",
    ]
    widths = [5, 11, 34, 18, 10, 7, 10, 11, 7, 15, 30, 10, 48, 40, 44, 44, 34, 40, 34, 26, 14, 14, 12, 28]
    if run_extras:
        reg_headers = reg_headers + ["Deep-verified", "Fix status", "Tests after fix"]
        widths = widths + [12, 30, 34]
    reg_rows = []
    for f in findings:
        dupes = "\n".join(
            f"{d.get('file')}:{d.get('line_start')}" for d in (f.get("duplicates") or [])
        )
        source_sink = " → ".join(x for x in (f.get("source_ref"), f.get("sink_ref")) if x)
        lines = f"{f.get('line_start')}-{f.get('line_end')}"
        pre = f.get("preconditions") or []
        row = [
            f.get("idx"), (f.get("severity") or "").title(), f.get("title"),
            f.get("vuln_class_label") or f.get("vuln_class"), f.get("cwe"),
            f.get("cvss_score"), f.get("cvss_rating"), f.get("confidence"),
            f.get("votes"), f.get("verdict"), f.get("file"), lines,
            _rich(f.get("description")), _rich(f.get("impact")),
            _rich(f.get("recommendation"), fix=True), _rich(f.get("exploit_scenario")),
            _rich("\n".join("• " + p for p in pre), bullets=False) if pre else "",
            f.get("code_snippet"), source_sink, dupes,
            "", "Open", "", "",
        ]
        if run_extras:
            fix = f.get("fix") or {}
            tests_label = FIX_TESTS_LABELS.get(fix.get("tests"), "")
            if fix.get("tests") == "broke_tests" and fix.get("tests_detail"):
                tests_label = f"{tests_label}: {fix['tests_detail']}"
            row += [
                "Yes" if f.get("deep") else "No",
                FIX_STATUS_LABELS.get(fix.get("status"), ""),
                tests_label,
            ]
        reg_rows.append(row)
    ws_reg = sheet("Findings Register", reg_headers, reg_rows, widths, tab="B02830")
    text_cols = {13, 14, 15, 16, 17, 18}
    for row in ws_reg.iter_rows(min_row=2):
        idx, sev, title, typ, cwe, cvss, rating, conf, votes, verdict, file_, lines = row[:12]
        for c in (idx, cwe, cvss, rating, conf, votes, lines):
            c.alignment = center
        title.font = bold_ink
        colors = _SEVERITY_FILLS.get(str(sev.value or "").lower())
        if colors:
            chip(sev, colors)
        if isinstance(cvss.value, (int, float)):
            cvss.number_format = "0.0"
            cvss.font = Font(bold=True, size=10, color=_SEVERITY_FILLS.get(str(sev.value or "").lower(), ("", INK))[1])
        if isinstance(conf.value, (int, float)):
            conf.number_format = "0%"
        if verdict.value in _VERDICT:
            label, bg, fg = _VERDICT[verdict.value]
            chip(verdict, (bg, fg), text=label, bold=False)
        else:
            verdict.value = ""
        file_.font = Font(name="Consolas", size=9.5, color=INK)
        row[17].font = Font(name="Consolas", size=9, color=INK)  # code
        row[18].font = Font(name="Consolas", size=9.5, color=INK)  # source -> sink
        row[19].font = Font(name="Consolas", size=9.5, color=MUTED)  # also at
        row[21].alignment = center
        chip(row[21], _STATUS_FILLS["Open"], bold=False)
        if run_extras:
            deep_cell, fix_cell, tests_cell = row[24], row[25], row[26]
            deep_cell.alignment = center
            broke = "Fix broke a test" in str(tests_cell.value or "")
            if broke:
                warn = _SEVERITY_FILLS["critical"]
                chip(fix_cell, warn, bold=True)
                chip(tests_cell, warn, bold=True)
                tests_cell.alignment = top_left
        # row height from the longest wrapped text cell (~ chars per line at the column width)
        longest = 1
        for col in text_cols:
            cell = row[col - 1]
            chars = _plain_len(cell.value)
            per_line = max(20, int(widths[col - 1] * 1.1))
            est = sum(max(1, -(-len(line) // per_line)) for line in str(cell.value if isinstance(cell.value, str) else "").split("\n")) if isinstance(cell.value, str) else max(1, -(-chars // per_line) + 1)
            longest = max(longest, est)
        ws_reg.row_dimensions[row[0].row].height = min(409, 14 * longest + 6)
    if reg_rows:
        last = len(reg_rows) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(_STATUS_OPTIONS) + '"', allow_blank=True)
        dv.prompt = "Pick the current state of this finding"
        ws_reg.add_data_validation(dv)
        dv.add(f"V2:V{last}")
        for option, (bg, fg) in _STATUS_FILLS.items():
            ws_reg.conditional_formatting.add(
                f"V2:V{last}",
                CellIsRule(operator="equal", formula=[f'"{option}"'], fill=fill(bg), font=Font(color=fg, size=10)),
            )
        ws_reg.freeze_panes = "D2"

    # -------------------------------------------------------- Exploit Chains
    ws_ch = wb.create_sheet("Exploit Chains")
    ws_ch.sheet_properties.tabColor = "C73E65"
    ws_ch.sheet_view.showGridLines = False
    intro = (
        "What is an exploit chain?  A chain is a set of findings an attacker would use one after "
        "another to get from a small foothold to real damage. Fixing any single step breaks the "
        "whole chain, so start with the step marked 'Fix this first'. Step numbers are the '#' "
        "column of the Findings Register."
    )
    ws_ch["A1"] = intro
    ws_ch.merge_cells("A1:F1")
    ws_ch["A1"].alignment = top_left
    ws_ch["A1"].font = Font(size=10, color=INK, italic=True)
    ws_ch["A1"].fill = fill(SECTION)
    ws_ch.row_dimensions[1].height = 48
    ch_headers = ["#", "Chain", "Severity", "Steps (in attack order)", "How the attack unfolds", "Fix this first"]
    ws_ch.append([_guard(h) for h in ch_headers])
    ws_ch.row_dimensions[2].height = 24
    for cell in ws_ch[2]:
        cell.font = header_font
        cell.fill = fill(BRAND)
        cell.alignment = center
        cell.border = cell_border
    ch_widths = [5, 36, 11, 46, 70, 40]
    for i, w in enumerate(ch_widths, start=1):
        ws_ch.column_dimensions[get_column_letter(i)].width = w
    if chains:
        for n, c in enumerate(chains, start=1):
            steps = [s for s in (c.get("steps") or [])]
            step_lines = "\n".join(
                f"{i}. #{s}  {title_by_idx.get(s) or 'finding not in register'}"
                for i, s in enumerate(steps, start=1)
            )
            first = steps[0] if steps else None
            fix_first = f"#{first}  {title_by_idx.get(first) or ''}".strip() if first is not None else ""
            ws_ch.append([n, _guard(c.get("title")), (c.get("severity") or "").title(),
                          _guard(step_lines), _rich(c.get("narrative")), _guard(fix_first)])
            r = ws_ch.max_row
            row = ws_ch[r]
            for cell in row:
                cell.alignment = top_left
                cell.font = body_font
                cell.border = cell_border
                if r % 2 == 1:
                    cell.fill = fill(ZEBRA)
            row[0].alignment = center
            row[1].font = bold_ink
            colors = _SEVERITY_FILLS.get(str(row[2].value or "").lower())
            if colors:
                chip(row[2], colors)
            row[5].font = Font(bold=True, size=10, color="0F766E")
            est = max(len(steps), -(-_plain_len(row[4].value) // 70) + 1)
            ws_ch.row_dimensions[r].height = min(409, 14 * est + 6)
    else:
        ws_ch.append(["", "No exploit chains reported", "", "", "The scanner did not link any findings together.", ""])
    ws_ch.freeze_panes = "A3"

    # ------------------------------------------------------------------ Fixes
    if fix_findings:
        fx_headers = ["#", "Title", "Severity", "Fix status", "Tests after fix",
                      "Files changed", "Root cause", "Remaining risks", "Patch"]
        fx_widths = [5, 32, 11, 30, 30, 30, 40, 40, 60]
        fx_rows = []
        for f in fix_findings:
            fix = f.get("fix") or {}
            deep = f.get("deep") or {}
            tests_label = FIX_TESTS_LABELS.get(fix.get("tests"), "")
            if fix.get("tests") == "broke_tests" and fix.get("tests_detail"):
                tests_label = f"{tests_label}: {fix['tests_detail']}"
            patch = fix.get("patch") or ""
            if len(patch) > 32000:
                patch = patch[:32000] + "... (truncated)"
            fx_rows.append([
                f.get("idx"), f.get("title"), (f.get("severity") or "").title(),
                FIX_STATUS_LABELS.get(fix.get("status"), fix.get("status") or ""),
                tests_label,
                "\n".join(fix.get("files") or []),
                deep.get("root_cause") or "",
                "\n".join(deep.get("remaining_risks") or []),
                patch,
            ])
        ws_fx = sheet("Fixes", fx_headers, fx_rows, fx_widths, tab=ACCENT)
        for row in ws_fx.iter_rows(min_row=2):
            idx, title, sev, status, tests_after, files_, root_cause, risks, patch_cell = row
            idx.alignment = center
            title.font = bold_ink
            colors = _SEVERITY_FILLS.get(str(sev.value or "").lower())
            if colors:
                chip(sev, colors)
            if "Fix broke a test" in str(tests_after.value or ""):
                warn = _SEVERITY_FILLS["critical"]
                chip(status, warn, bold=True)
                chip(tests_after, warn, bold=True)
                tests_after.alignment = top_left
            patch_cell.font = Font(name="Consolas", size=8.5, color=INK)
            patch_cell.alignment = top_left
            longest = max(1, str(patch_cell.value or "").count("\n") + 1, -(-_plain_len(patch_cell.value) // 90))
            ws_fx.row_dimensions[row[0].row].height = min(409, 14 * longest + 6)
        ws_fx.freeze_panes = "A2"

    # ------------------------------------------------------------ Scan Coverage
    if run_extras:
        cov = run_extras.get("coverage") or {}
        tests = run_extras.get("tests") or {}
        cov_rows = [
            ["Deep-verified run", ""],
            [f"{run_extras.get('deep_verified', 0)} of {run_extras.get('total', len(findings))} findings were "
             "deep-verified by the scanner (its top findings by CVSS); the rest come from the first review "
             "pass only.", ""],
        ]
        if cov:
            cov_rows.append([])
            cov_rows.append(["Coverage metrics", ""])
            for label, key in (
                ("Coverage %", "coverage_pct"), ("Files in scope", "files_in_scope"),
                ("Files analyzed", "files_analyzed"), ("Chunks", "chunks"),
            ):
                if cov.get(key) is not None:
                    cov_rows.append([label, cov[key]])
            if cov.get("health"):
                cov_rows.append([])
                cov_rows.append(["Scan health", ""])
                for line in cov["health"]:
                    cov_rows.append([line, ""])
            if cov.get("threat_model"):
                cov_rows.append([])
                cov_rows.append(["Threat model", ""])
                cov_rows.append([cov["threat_model"], ""])
        if tests:
            cov_rows.append([])
            cov_rows.append(["Test summary", ""])
            for label, key in (
                ("Build", "build"), ("Suites", "suites"), ("Cases", "cases"),
                ("Failures", "failures"), ("Errors", "errors"), ("Skipped", "skipped"),
            ):
                if tests.get(key) is not None:
                    cov_rows.append([label, tests[key]])
            for fc in tests.get("failing") or []:
                cov_rows.append([f"Failing: {fc.get('test')}", fc.get("message") or ""])
            if tests.get("not_run_modules"):
                cov_rows.append(["Modules not tested", ", ".join(tests["not_run_modules"])])
        ws_cov = sheet("Scan Coverage", ["Item", "Detail"], cov_rows, [40, 70], filters=False, tab=ACCENT)
        section_labels = ("Coverage metrics", "Scan health", "Threat model", "Test summary")
        for row in ws_cov.iter_rows(min_row=2, max_col=2):
            label, _detail = row
            key = str(label.value or "")
            if key in section_labels or key == "Deep-verified run":
                for c in row:
                    c.fill = fill(SECTION)
                    c.font = Font(bold=True, size=10.5, color=BRAND)

    wb.properties.title = f"{review.name} - Code Security Review"
    wb.properties.creator = branding.get("report_display_name") or "ScopeSense"
    # openpyxl stamps "now" into core.xml, which broke byte-level determinism
    # whenever two builds straddled a second boundary; pin it.
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)

    buf = io.BytesIO()
    wb.save(buf)
    return _normalize_zip(buf.getvalue())


def _normalize_zip(data: bytes) -> bytes:
    """openpyxl stamps every zip entry with 'now' (2-second resolution), so two
    builds a moment apart differ byte-for-byte. Rewrite entries with a fixed
    timestamp so the same review always yields identical bytes."""
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            dst.writestr(fixed, src.read(info.filename))
    return out.getvalue()
