"""Ingest the SME's returned LEGAL_SEVERITY_SME_PACK_*.xlsx and print what
LEGAL_SEVERITY_CALIBRATION.md section 4 asks for: agreement rate per finding
type and the systematic over/under-severity patterns.

    python scripts/severity_calibration_ingest.py docs/planning/severity_calibration/LEGAL_SEVERITY_SME_PACK_2026-09-12.xlsx
    python scripts/severity_calibration_ingest.py --self-check

Reads the Findings sheet by header name, so column order changes are tolerated;
rows with no SME severity are reported as "unrated" and excluded from rates.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

RANK = {"critical": 4, "major": 3, "medium": 2, "low": 1, "info": 0}


def read_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["Findings"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for r in rows:
        if not r or r[0] is None:
            continue
        out.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    return out


def analyse(rows: list[dict]) -> dict:
    per_type: dict[str, dict] = defaultdict(lambda: {"rated": 0, "agree": 0, "over": 0, "under": 0, "not_finding": 0, "delta_sum": 0})
    unrated = 0
    for r in rows:
        t = str(r.get("Finding type") or "?")
        sys_sev = str(r.get("System severity") or "").strip().lower()
        sme = str(r.get("SME severity") or "").strip().lower()
        agree = str(r.get("Agree?") or "").strip().lower()
        if agree == "not a finding":
            per_type[t]["not_finding"] += 1
            per_type[t]["rated"] += 1
            continue
        if sme not in RANK or sys_sev not in RANK:
            unrated += 1
            continue
        d = per_type[t]
        d["rated"] += 1
        delta = RANK[sys_sev] - RANK[sme]  # >0 system rated higher than SME
        d["delta_sum"] += delta
        if delta == 0:
            d["agree"] += 1
        elif delta > 0:
            d["over"] += 1
        else:
            d["under"] += 1
    return {"per_type": dict(per_type), "unrated": unrated, "total": len(rows)}


def report(a: dict) -> str:
    lines = [f"{'Finding type':<34}{'rated':>6}{'agree%':>8}{'over':>6}{'under':>6}{'notF':>6}{'mean d':>8}  pattern"]
    tot = {"rated": 0, "agree": 0, "over": 0, "under": 0}
    for t, d in sorted(a["per_type"].items()):
        scored = d["rated"] - d["not_finding"]
        rate = (d["agree"] / scored * 100) if scored else 0
        mean = (d["delta_sum"] / scored) if scored else 0
        if scored == 0:
            pat = "all 'not a finding'" if d["not_finding"] else "no ratings"
        elif d["over"] >= max(2, 0.5 * scored):
            pat = "SYSTEMATICALLY OVER-SEVERE"
        elif d["under"] >= max(2, 0.5 * scored):
            pat = "SYSTEMATICALLY UNDER-SEVERE"
        elif rate >= 80:
            pat = "calibrated"
        else:
            pat = "mixed"
        lines.append(f"{t:<34}{d['rated']:>6}{rate:>7.0f}%{d['over']:>6}{d['under']:>6}{d['not_finding']:>6}{mean:>+8.2f}  {pat}")
        for k in tot:
            tot[k] += d[k]
    scored = sum(d["rated"] - d["not_finding"] for d in a["per_type"].values())
    overall = (tot["agree"] / scored * 100) if scored else 0
    lines.append(f"\noverall agreement {overall:.0f}% on {scored} rated findings; over {tot['over']}, under {tot['under']}, unrated {a['unrated']} of {a['total']}")
    lines.append("mean d: +1.00 = system one tier higher than the SME on average; -1.00 = one tier lower.")
    return "\n".join(lines)


def self_check() -> None:
    rows = [
        {"Finding type": "missing_liability_cap", "System severity": "major", "SME severity": "major", "Agree?": "Yes"},
        {"Finding type": "missing_liability_cap", "System severity": "major", "SME severity": "critical", "Agree?": "No"},
        {"Finding type": "missing_warranty", "System severity": "major", "SME severity": "low", "Agree?": "No"},
        {"Finding type": "missing_warranty", "System severity": "major", "SME severity": "medium", "Agree?": "No"},
        {"Finding type": "ambiguous_legal_language", "System severity": "medium", "SME severity": "", "Agree?": ""},
        {"Finding type": "no_governing_law", "System severity": "major", "SME severity": "", "Agree?": "Not a finding"},
    ]
    a = analyse(rows)
    assert a["unrated"] == 1
    assert a["per_type"]["missing_liability_cap"]["agree"] == 1 and a["per_type"]["missing_liability_cap"]["under"] == 1
    assert a["per_type"]["missing_warranty"]["over"] == 2 and a["per_type"]["missing_warranty"]["delta_sum"] == 3
    assert a["per_type"]["no_governing_law"]["not_finding"] == 1
    out = report(a)
    assert "SYSTEMATICALLY OVER-SEVERE" in out and "missing_warranty" in out
    print("self-check ok")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--self-check":
        self_check()
    elif len(sys.argv) > 1:
        print(report(analyse(read_rows(Path(sys.argv[1])))))
    else:
        sys.exit(__doc__)
