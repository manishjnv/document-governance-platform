"""Generate the /resources/templates downloads into apps/web/public/templates/.

Two files, both deterministic and dependency-light (reportlab + openpyxl):
  sow-review-checklist.pdf                  -- the 10-point pre-signature checklist
  rfp-evaluation-criteria-worksheet.xlsx    -- weighted scoring sheet, FAR Part 15 shape
The MITRE environment template is produced by scripts/generate_mitre_templates.py
(unchanged). Run from the repo root: python scripts/generate_templates.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parents[1] / "apps" / "web" / "public" / "templates"
BRAND = "0057B8"

CHECKLIST = [
    ("Scope and deliverables", [
        ("Deliverables are concrete and testable", "\"A responsive site with these 12 named pages\", not \"a modern website\". Vague deliverable language is the largest single source of scope creep."),
        ("Every deliverable has acceptance criteria", "A stated process for how it is reviewed and signed off. Without one, \"done\" is a matter of opinion."),
        ("Exclusions are stated explicitly", "What is not included is written down, not inferred from what is."),
    ]),
    ("Timeline and process", [
        ("A change-control clause exists", "Any scope addition goes through a written, priced amendment before work starts on it."),
        ("Milestones and dependencies are dated and owned", "Who delivers what, by when, is unambiguous."),
        ("Assumptions and client-side dependencies are listed", "Approvals, access and third-party data have owners and dates; a timeline that assumes instant client turnaround rarely survives."),
    ]),
    ("Commercial and legal terms", [
        ("A liability cap with a defined amount and carve-out list", "Or liability is explicitly unbounded and someone has accepted that."),
        ("Payment milestones are tied to deliverables", "Not calendar dates alone, so payment and delivery stay linked."),
        ("The MSA is referenced correctly", "The SOW does not quietly restate or contradict terms the MSA already covers; the order-of-precedence clause is known."),
    ]),
    ("The final check", [
        ("Read it as the other side would", "A SOW written entirely from one party's perspective has gaps the other party will surface later, during the project rather than before it."),
    ]),
]


def build_pdf() -> Path:
    path = OUT / "sow-review-checklist.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
        title="SOW review checklist", author="ScopeWise", subject="Ten-point pre-signature Statement of Work checklist",
    )
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=20, alignment=0, spaceAfter=4, textColor="#0F172A")
    sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=10, textColor="#3F4A5C", spaceAfter=14)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=13, textColor="#" + BRAND, spaceBefore=10, spaceAfter=4)
    item = ParagraphStyle("item", parent=ss["Normal"], fontSize=10.5, leading=14, leftIndent=18, spaceAfter=2)
    note = ParagraphStyle("note", parent=ss["Normal"], fontSize=9.5, leading=12.5, leftIndent=18, textColor="#3F4A5C", spaceAfter=7)
    foot = ParagraphStyle("foot", parent=ss["Normal"], fontSize=8.5, textColor="#3F4A5C", spaceBefore=16)

    story = [
        Paragraph("SOW review checklist", h1),
        Paragraph("Ten places risk most often hides in a Statement of Work. Tick each one before signature; anything unticked is a finding to raise.", sub),
    ]
    n = 0
    for section, items in CHECKLIST:
        story.append(Paragraph(section, h2))
        for title, why in items:
            n += 1
            story.append(Paragraph(f"[&nbsp;&nbsp;]&nbsp; <b>{n}. {title}</b>", item))
            story.append(Paragraph(why, note))
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        "ScopeWise runs this checklist, plus a deterministic rule engine and six specialist reviewers, on every uploaded SOW or RFP and quotes the clause behind each finding. "
        "https://scopewise.assessiq.in/product/sow-review &nbsp;&middot;&nbsp; This checklist is guidance, not legal advice.", foot))
    doc.build(story)
    return path


FACTORS = [
    ("Technical approach", 30, "Does the proposal meet each stated requirement? Cite requirement IDs."),
    ("Past performance", 20, "Comparable engagements in the last three years, with references."),
    ("Transition and delivery plan", 15, "Milestones, dependencies on the buyer, cutover and rollback."),
    ("Team and key personnel", 10, "Named roles, availability, substitution terms."),
    ("Price", 25, "Total evaluated price in the required format; note assumptions and exclusions."),
]


def build_xlsx() -> Path:
    path = OUT / "rfp-evaluation-criteria-worksheet.xlsx"
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=BRAND)

    def sheet(ws, headers, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(vertical="center", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 100
    for line in (
        "RFP evaluation criteria worksheet (ScopeWise)",
        "",
        "Structure mirrors FAR Part 15 (contracting by negotiation): requirements, evaluation factors and weights, submission instructions, basis for award.",
        "1. Requirements: one row per requirement with an ID, so proposals can be scored against the same yardstick.",
        "2. Evaluation factors: weights must sum to 100. State them in the RFP itself, not only here.",
        "3. Scoring: one column per bidder; score 0-5 per factor; the weighted total is computed for you.",
        "4. Basis for award: fill in the sentence that goes into the RFP (for example, highest weighted score, or lowest price technically acceptable).",
        "",
        "This worksheet is guidance, not legal advice. https://scopewise.assessiq.in/use-cases/rfp-review",
    ):
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)

    req = wb.create_sheet("Requirements")
    sheet(req, ["Req ID", "Requirement (testable statement)", "Mandatory / Desirable", "Section of RFP", "How compliance is verified"], [10, 60, 20, 16, 40])
    for i in range(1, 4):
        req.append([f"R-{i:03d}", "", "Mandatory", "", ""])

    ev = wb.create_sheet("Evaluation Factors")
    sheet(ev, ["Factor", "Weight (%)", "What the evaluator looks for", "Bidder A (0-5)", "Bidder B (0-5)", "Bidder C (0-5)", "Weighted A", "Weighted B", "Weighted C"], [30, 12, 55, 14, 14, 14, 12, 12, 12])
    for r, (name, w, what) in enumerate(FACTORS, start=2):
        ev.append([name, w, what, None, None, None, f"=B{r}*D{r}/5", f"=B{r}*E{r}/5", f"=B{r}*F{r}/5"])
    last = len(FACTORS) + 1
    ev.append(["Total", f"=SUM(B2:B{last})", "Weights must sum to 100", None, None, None, f"=SUM(G2:G{last})", f"=SUM(H2:H{last})", f"=SUM(I2:I{last})"])
    for c in ev[last + 1]:
        c.font = Font(bold=True)

    sub = wb.create_sheet("Submission & Award")
    sheet(sub, ["Item", "Your RFP says"], [34, 90])
    for row in (
        ("Submission deadline (date, time, time zone)", ""),
        ("Q&A window and how answers are shared with all bidders", ""),
        ("Required proposal format and page limits", ""),
        ("Required pricing format", ""),
        ("Mandatory terms (state in the body, not only an attachment)", ""),
        ("Basis for award (one sentence)", "Award to the responsive proposal with the highest weighted score."),
    ):
        sub.append(list(row))

    wb.properties.creator = "ScopeWise"
    wb.properties.title = "RFP evaluation criteria worksheet"
    wb.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for p in (build_pdf(), build_xlsx()):
        print(f"wrote {p} ({p.stat().st_size:,} bytes)")
